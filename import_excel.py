#!/usr/bin/env python3
"""Import Excel/CSV package/version rows, verify live CNB data, and rebuild the website.

Usage: python import_excel.py packages.xlsx [more.xlsx] [--dry-run]
Requires Python 3.10+ (plus openpyxl for Excel). No CNB credentials are needed.
"""

import argparse
import csv
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from hashlib import sha256
import json
import io
from math import ceil
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

from sync_cnb import NOTE, key, read_snapshot

ROOT = Path(__file__).resolve().parent
CNB_URL = "https://cnb.cool/OpenHarmonyPCDeveloper/pypi"
ALIASES = {
    "name": {"库名", "三方库名称", "三方库名", "包名", "包名称", "名称", "name", "package", "package_name", "library", "library_name"},
    "version": {"版本号", "版本", "version", "package_version", "library_version"},
    "date": {"最终成功日期", "成功日期", "完成日期", "适配完成日期", "completed_at", "completion_date", "date"},
}


def now():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def log(message):
    print(message, flush=True)


class CnbClient:
    def __init__(self, workers=6, timeout=30, retries=3):
        self.workers, self.timeout, self.retries = workers, timeout, retries

    def request(self, path, **params):
        url = CNB_URL + path + "?" + urllib.parse.urlencode(params)
        headers = {"User-Agent": "ohos-python-support-list/1.0", "Accept": "application/vnd.cnb.api+json"}
        for attempt in range(self.retries):
            try:
                with urllib.request.urlopen(urllib.request.Request(url, headers=headers), timeout=self.timeout) as response:
                    data = json.load(response)
                    response_headers = {k.lower(): v for k, v in response.headers.items()}
                return data, response_headers
            except (urllib.error.URLError, OSError, ValueError) as exc:
                if isinstance(exc, urllib.error.HTTPError) and exc.code not in {408, 429, 500, 502, 503, 504}:
                    raise RuntimeError(f"CNB 请求失败（HTTP {exc.code}）：{url}") from exc
                if attempt + 1 == self.retries:
                    raise RuntimeError(f"CNB 读取失败，不能判为不存在：{url}；{exc}") from exc
                time.sleep(min(2 ** attempt, 8))

    @staticmethod
    def pagination(headers, page):
        try:
            total, size, actual_page = (int(headers[k]) for k in ("x-cnb-total", "x-cnb-page-size", "x-cnb-page"))
        except (KeyError, ValueError) as exc:
            raise ValueError("CNB 未返回完整分页信息") from exc
        if total < 0 or size <= 0 or actual_page != page:
            raise ValueError("CNB 分页信息异常")
        return total, size

    def inventory(self):
        def page(number):
            data, headers = self.request("/-/packages", type="pypi", page=number, page_size=100, ordering="name_ascend")
            total, size = self.pagination(headers, number)
            if not isinstance(data, list):
                raise ValueError("CNB 制品列表格式异常")
            return number, data, total, size

        _, first, total, size = page(1)
        if not total:
            raise ValueError("CNB 返回空仓库，停止更新，请核查服务状态")
        count = ceil(total / size)
        pages = {1: first}
        with ThreadPoolExecutor(max_workers=self.workers) as pool:
            for future in as_completed([pool.submit(page, n) for n in range(2, count + 2)]):
                number, data, current_total, current_size = future.result()
                if (current_total, current_size) != (total, size):
                    raise ValueError("CNB 列表在分页读取期间发生变化")
                pages[number] = data
        if pages[count + 1]:
            raise ValueError("CNB 末页之外仍有记录")
        items = [item for n in range(1, count + 1) for item in pages[n]]
        if len(items) != total or len({key(x["name"], "")[0] for x in items}) != total:
            raise ValueError("CNB 库列表存在重复或漏页")
        packages = [{"name": x["name"], "version_count": int(x["count"]),
                     "latest_version": x["last_artifact_name"],
                     "last_push_at": x.get("last_pusher", {}).get("push_at", "")} for x in items]
        if any(p["version_count"] < 1 for p in packages):
            raise ValueError("CNB 库存在异常版本计数")
        return sorted(packages, key=lambda p: p["name"]), {
            "pages": count, "page_size": size, "reported_packages": total, "empty_end_page_checked": True,
        }

    def versions(self, package):
        name, expected = package["name"], package["version_count"]
        if expected == 1 and package["latest_version"]:
            return [{"name": name, "version": package["latest_version"], "last_push_at": package["last_push_at"]}]
        result, number = [], 1
        while True:
            data, headers = self.request("/-/packages/pypi/" + urllib.parse.quote(name, safe="") + "/-/tags",
                                         page=number, page_size=100)
            total, size = self.pagination(headers, number)
            tags = data.get("pypi")
            if total != expected or not isinstance(tags, list):
                raise ValueError(f"{name} 的版本数量或格式与库列表不一致")
            result.extend({"name": name, "version": t["name"], "last_push_at": t.get("last_pusher", {}).get("push_at", "")} for t in tags)
            if number * size >= total:
                break
            if not tags:
                raise ValueError(f"{name} 的版本分页提前结束")
            number += 1
        if len(result) != expected or len({v["version"] for v in result}) != expected:
            raise ValueError(f"{name} 的版本记录不完整或重复")
        return result

    def snapshot(self):
        # A push can move records between pages. Require a stable full second pass.
        for attempt in range(2):
            started = now()
            try:
                packages, pagination = self.inventory()
                log(f"CNB：{len(packages):,} 个库，正在核验全部版本……")
                versions = []
                with ThreadPoolExecutor(max_workers=self.workers) as pool:
                    for future in as_completed([pool.submit(self.versions, p) for p in packages]):
                        versions.extend(future.result())
                final_packages, _ = self.inventory()
                if packages != final_packages:
                    raise ValueError("CNB 在核验过程中有更新")
                expected = sum(p["version_count"] for p in packages)
                if len(versions) != expected or len({key(v["name"], v["version"]) for v in versions}) != expected:
                    raise ValueError("CNB 版本总量或唯一性校验失败")
                return {"metadata": {
                    "started_at": started, "finished_at": now(), "source": CNB_URL,
                    "packages": len(packages), "versions": len(versions), "expected_versions": expected,
                    "pagination": pagination, "inventory_stable_across_two_full_passes": True, "failures": [],
                    "method": "All package pages; sole version when count=1; all tag pages otherwise; stable second full inventory.",
                }, "packages": packages, "versions": sorted(versions, key=lambda v: key(v["name"], v["version"]))}
            except ValueError:
                if attempt:
                    raise
                log("CNB 数据发生变化或计数不一致，重新完整核验一次……")
        raise RuntimeError("CNB 核验未完成")


def header_key(value):
    return re.sub(r"[\s_()（）-]+", "", str(value or "").casefold())


def date_text(value):
    if value is None or str(value).strip() in {"", "-", "—"}:
        return "-"
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    match = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?(?:[ T]\d{2}:\d{2}:\d{2})?", str(value).strip())
    if not match:
        raise ValueError("成功日期应为 Excel 日期或 YYYY-MM-DD")
    return date(*map(int, match.groups())).isoformat()


def input_tables(path, sheets=None):
    """Yield table names and raw values; never coerce version identifiers to numbers."""
    if path.suffix.lower() == ".csv":
        if sheets:
            raise ValueError("CSV 不支持 --sheet，请单独导入或去掉该参数")
        raw = path.read_bytes()
        try:
            content = raw.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            content = raw.decode("gb18030")
            encoding = "gb18030"
        with io.StringIO(content, newline="") as stream:
            yield "CSV", csv.reader(stream, strict=True), encoding
        return
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"仅支持 .xlsx / .xlsm / .csv，请先转换文件：{path}")
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先运行：python -m pip install -r requirements.txt") from exc
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheets and set(sheets) - set(book.sheetnames):
            raise ValueError(f"{path.name} 中不存在工作表：{sorted(set(sheets) - set(book.sheetnames))}")
        for sheet in book:
            if not sheets or sheet.title in sheets:
                yield sheet.title, sheet.iter_rows(values_only=True), None
    finally:
        book.close()


def read_excel(paths, sheets=None, name_column=None, version_column=None, date_column=None):
    """Read Excel or CSV inputs (name retained for existing callers)."""
    overrides = {"name": name_column, "version": version_column, "date": date_column}
    aliases = {field: {header_key(overrides[field])} if overrides[field] else {header_key(v) for v in values}
               for field, values in ALIASES.items()}
    accepted, errors, sheet_report = [], [], []
    for path in paths:
        tables = input_tables(path, sheets)
        try:
            for title, rows, encoding in tables:
                columns = None
                count = 0
                for number, values in enumerate(rows, 1):
                    if columns is None:
                        if number > 30:
                            break
                        found = {field: [i for i, value in enumerate(values) if header_key(value) in names]
                                 for field, names in aliases.items()}
                        if len(found["name"]) == 1 and len(found["version"]) == 1:
                            if len(found["date"]) > 1:
                                raise ValueError(f"{title} 有多个日期列，请通过 --date-column 指定")
                            columns = {field: indexes[0] for field, indexes in found.items() if indexes}
                            if date_column and "date" not in columns:
                                raise ValueError(f"{title} 缺少指定的日期列：{date_column}")
                        continue
                    def cell(field):
                        index = columns.get(field)
                        return values[index] if index is not None and index < len(values) else None
                    name, version = cell("name"), cell("version")
                    if all(value is None or (isinstance(value, str) and not value.strip()) for value in (name, version)):
                        continue
                    if header_key(name) in aliases["name"] and header_key(version) in aliases["version"]:
                        continue
                    location = {"file": str(path), "sheet": title, "row": number}
                    try:
                        if not isinstance(name, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", name.strip()):
                            raise ValueError("库名为空、为公式或包含不支持的字符")
                        if not isinstance(version, str):
                            raise ValueError("版本号必须以文本保存；数值单元格可能丢失 1.10 等版本的末尾零，请重新输入原始文本")
                        if version.strip() in {"", "-"} or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._!+\-]*", version.strip()):
                            raise ValueError("版本号为空、为公式或非单个精确版本号")
                        accepted.append({"name": name.strip(), "version": version.strip(),
                                         "completed_at": date_text(cell("date")), **location})
                        count += 1
                    except ValueError as exc:
                        errors.append({**location, "name": str(name), "version": str(version), "error": str(exc)})
                if columns is None and (sheets or encoding):
                    raise ValueError(f"{path.name} 工作表 {title} 前 30 行未找到库名和版本号表头")
                sheet_report.append({"file": str(path), "sheet": title, "encoding": encoding,
                                     "data_sheet": columns is not None, "valid_rows": count})
        finally:
            tables.close()
    return accepted, errors, sheet_report


def merge_source(original, imported, cnb):
    retained, prefix, duplicates, changes = {}, [], [], []
    before = 0
    for line in original.decode("utf-8-sig").splitlines():
        if not re.match(r"^\|\s*\d+\s*\|", line):
            if not retained and line != NOTE:
                prefix.append(line)
            continue
        row = [x.strip() for x in line.strip().strip("|").split("|")]
        if len(row) != 5 or row[3] not in {"是", "否"}:
            raise ValueError(f"支持列表数据行格式错误：{line}")
        before += 1
        pair = key(row[1], row[2])
        status = pair in cnb
        if status != (row[3] == "是"):
            changes.append({"name": row[1], "version": row[2], "before": row[3], "after": "是" if status else "否"})
        if pair in retained:
            duplicates.append({"name": row[1], "version": row[2], "original_id": row[0]})
            retained[pair]["completed_at"] = max(retained[pair]["completed_at"], row[4])
        else:
            retained[pair] = {"name": row[1], "version": row[2], "completed_at": row[4], "adapted": status}
    if not retained:
        raise ValueError("现有 support_list.md 为空或没有合法数据行")
    imported_keys, excel_added, date_updates = set(), [], []
    for item in imported:
        pair = key(item["name"], item["version"])
        imported_keys.add(pair)
        if pair not in retained:
            retained[pair] = {field: item[field] for field in ("name", "version", "completed_at")}
            retained[pair]["adapted"] = pair in cnb
            excel_added.append(pair)
        elif item["completed_at"] != "-" and item["completed_at"] > retained[pair]["completed_at"]:
            date_updates.append({"name": item["name"], "version": item["version"],
                                 "before": retained[pair]["completed_at"], "after": item["completed_at"]})
            retained[pair]["completed_at"] = item["completed_at"]
    cnb_added = sorted(cnb.keys() - retained.keys())
    for pair in cnb_added:
        retained[pair] = {"name": cnb[pair]["name"], "version": cnb[pair]["version"], "completed_at": "-", "adapted": True}
    assert {pair for pair, record in retained.items() if record["adapted"]} == set(cnb)
    while len(prefix) > 1 and not prefix[1].strip():
        prefix.pop(1)
    prefix[1:1] = ["", NOTE, ""]
    rendered = [f"| {i} | {r['name']} | {r['version']} | {'是' if r['adapted'] else '否'} | {r['completed_at']} |"
                for i, r in enumerate(retained.values(), 1)]
    summary = {"existing_rows": before, "excel_rows": len(imported), "excel_unique_versions": len(imported_keys),
               "excel_duplicate_rows": len(imported) - len(imported_keys), "excel_added_versions": len(excel_added),
               "cnb_added_versions": len(cnb_added), "existing_duplicates_merged": len(duplicates),
               "existing_status_changes": len(changes), "date_updates": len(date_updates),
               "total": len(retained), "yes": len(cnb), "no": len(retained) - len(cnb)}
    names = {name for name, version in cnb}
    checked = [{**item, "status": "是" if key(item["name"], item["version"]) in cnb else "否",
                "reason": "CNB 版本存在" if key(item["name"], item["version"]) in cnb else
                          "CNB 中存在该库但无此版本" if key(item["name"], "")[0] in names else "CNB 未收录该库"}
               for item in imported]
    return ("\n".join(prefix + rendered) + "\n").encode("utf-8"), {
        "summary": summary, "excel_rows": checked, "status_changes": changes, "date_updates": date_updates,
        "duplicates_merged": duplicates, "excel_added": excel_added, "cnb_added": cnb_added,
    }


def verify_generated(stage, cnb):
    packages = json.loads((stage / "data/all.json").read_text(encoding="utf-8"))
    by_id = {p["id"]: p for p in packages}
    if len(by_id) != len(packages) or len({key(p["name"], p["version"]) for p in packages}) != len(packages):
        raise ValueError("生成数据存在重复编号或版本")
    if {key(p["name"], p["version"]) for p in packages if p["adapted"]} != set(cnb):
        raise ValueError("网页中标为“是”的版本与 CNB 不一致")
    for folder in ("pages", "search", "categories"):
        ids = []
        for path in (stage / "data" / folder).glob("*.json"):
            items = json.loads(path.read_text(encoding="utf-8"))
            if any(p != by_id.get(p["id"]) for p in items):
                raise ValueError(f"网页分片数据不一致：{path.name}")
            ids.extend(p["id"] for p in items)
        if Counter(ids) != Counter({i: 2 if folder == "pages" else 1 for i in by_id}):
            raise ValueError(f"网页 {folder} 分片总量不一致")


def build_stage(repo, run, source_bytes, verification, cnb):
    stage = run / "staged"
    stage.mkdir()
    for name in ("gen_packages.py", "README.md"):
        shutil.copy2(repo / name, stage / name)
    (stage / "support_list.md").write_bytes(source_bytes)
    write_json(stage / "cnb_verification.json", verification)
    subprocess.run([sys.executable, "-B", str(stage / "gen_packages.py")], cwd=stage, check=True)
    validator = repo / ".agents/skills/classify-ohos-python-packages/scripts/validate_categories.py"
    subprocess.run([sys.executable, "-B", str(validator), "--repo-root", str(stage)], check=True)
    verify_generated(stage, cnb)
    return stage


def publish_local(repo, stage, run, original):
    # Stage first; preserve the entire prior data directory, including any untracked files.
    for name in ("data", "support_list.md", "cnb_verification.json"):
        path = repo / name
        if path.is_symlink() or (path.exists() and path.resolve().parent != repo):
            raise ValueError(f"目标路径不在项目内或为链接，停止替换：{path}")
    if (repo / "support_list.md").read_bytes() != original:
        raise ValueError("核验期间 support_list.md 已被其他程序修改，请重新运行")
    backup = run / "backup"
    backup.mkdir()
    names = ("data", "support_list.md", "cnb_verification.json")
    saved, installed = [], []
    try:
        for name in names:
            if (repo / name).exists():
                (repo / name).rename(backup / name)
                saved.append(name)
            (stage / name).rename(repo / name)
            installed.append(name)
    except BaseException:
        for name in reversed(installed):
            (repo / name).rename(stage / name)
        for name in reversed(saved):
            (backup / name).rename(repo / name)
        raise
    return backup


def parser():
    result = argparse.ArgumentParser(description="读取 Excel/CSV，在线核验 CNB，合并支持列表并生成网页数据。")
    result.add_argument("excel", nargs="+", type=Path, help="一个或多个 .xlsx / .xlsm / .csv 文件")
    result.add_argument("--dry-run", action="store_true", help="在线核验并预览，只保存报告，不改支持列表或网页")
    result.add_argument("--sheet", action="append", help="仅 Excel：只读取指定工作表，可重复；默认识别全部数据表")
    result.add_argument("--name-column", help="自定义库名列标题")
    result.add_argument("--version-column", help="自定义版本列标题")
    result.add_argument("--date-column", help="自定义成功日期列标题")
    result.add_argument("--skip-invalid", action="store_true", help="明确跳过无效行并写入报告；默认发现无效行即停止")
    result.add_argument("--workers", type=int, choices=range(1, 13), default=6, metavar="1-12", help="CNB 并发数，默认 6")
    result.add_argument("--timeout", type=int, default=30, help="单次请求超时秒数，默认 30")
    result.add_argument("--retries", type=int, choices=range(1, 6), default=3, metavar="1-5", help="单次请求尝试次数，默认 3")
    result.add_argument("--repo-root", type=Path, default=ROOT, help="目标项目目录，默认脚本所在目录")
    return result


def run_import(args):
    repo = args.repo_root.resolve()
    paths = [p.resolve(strict=True) for p in args.excel]
    for name in ("support_list.md", "gen_packages.py", "README.md", ".agents/skills/classify-ohos-python-packages/scripts/validate_categories.py"):
        if not (repo / name).is_file():
            raise ValueError(f"目标项目缺少文件：{repo / name}")
    if args.timeout <= 0:
        raise ValueError("--timeout 必须大于 0")
    run = repo / "outputs" / "excel_import" / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run.mkdir(parents=True)
    report = {"started_at": now(), "mode": "preview" if args.dry_run else "apply", "inputs": [str(p) for p in paths]}
    try:
        imported, errors, sheets = read_excel(paths, args.sheet, args.name_column, args.version_column, args.date_column)
        report.update({"invalid_rows": errors, "sheets": sheets})
        if errors and not args.skip_invalid:
            raise ValueError(f"输入表格有 {len(errors)} 行不完整或格式异常，详情见 report.json；修正后再运行，或明确使用 --skip-invalid")
        if not imported:
            raise ValueError("输入表格中没有可导入的库名和版本记录")
        log(f"输入表格：{len(imported):,} 条有效记录，{len(errors)} 条异常；开始在线核验 CNB。")
        original = (repo / "support_list.md").read_bytes()
        snapshot = CnbClient(args.workers, args.timeout, args.retries).snapshot()
        snapshot_path = run / "cnb_snapshot.json"
        write_json(snapshot_path, snapshot)
        metadata, cnb = read_snapshot(snapshot_path)
        source_bytes, result = merge_source(original, imported, cnb)
        report.update(result)
        report["cnb_checked_at"] = metadata["finished_at"]
        report["source_before_sha256"] = sha256(original).hexdigest()
        report["source_after_sha256"] = sha256(source_bytes).hexdigest()
        verification = {"source_url": CNB_URL, "checked_at": metadata["finished_at"],
                        "package_count": metadata["packages"], "version_count": len(cnb),
                        "rule": "normalized-name-and-exact-version", "source_sha256": sha256(source_bytes).hexdigest(),
                        "snapshot_sha256": sha256(snapshot_path.read_bytes()).hexdigest()}
        if not args.dry_run:
            stage = build_stage(repo, run, source_bytes, verification, cnb)
            backup = publish_local(repo, stage, run, original)
            report["backup"] = str(backup)
        report["status"] = "previewed" if args.dry_run else "updated"
        log(json.dumps(report["summary"], ensure_ascii=False, indent=2))
        return run, report
    except BaseException as exc:
        report.update({"status": "failed", "error": str(exc)})
        raise
    finally:
        report["finished_at"] = now()
        write_json(run / "report.json", report)
        log(f"处理报告：{run / 'report.json'}")


def main():
    try:
        args = parser().parse_args()
        run, report = run_import(args)
        if args.dry_run:
            log("预览完成。去掉 --dry-run 即可更新 support_list.md 和网页数据。")
        else:
            log("更新完成，刷新本地网页即可查看。GitHub Pages 需提交并推送项目后发布。")
    except (Exception, KeyboardInterrupt) as exc:
        print(f"处理失败：{exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
